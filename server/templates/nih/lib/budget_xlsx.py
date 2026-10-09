# out/budget.xlsx. With budget-excel.yaml (an institution's layer) and its
# workbook: budget.yaml's inputs, written into a copy of that workbook, whose
# own formulas do the rest. Otherwise a plain workbook of the budget
# lib/budget.typ computes, its totals live formulas.
#   python3 lib/budget_xlsx.py          export
#   python3 lib/budget_xlsx.py import F read workbook F back into budget.yaml
import json, os, subprocess, sys
from openpyxl import Workbook, load_workbook
from ruamel.yaml import YAML

yaml = YAML()  # round-trip: budget.yaml keeps its comments and layout
yaml.indent(mapping=2, sequence=4, offset=2)
typst = os.environ.get('TYPST', 'typst')
computed = lambda: json.loads(subprocess.run([typst, 'eval', 'import "/lib/budget.typ": budget; budget', '--in', 'main.typ'],
                                             capture_output=True, text=True, check=True).stdout)
xl = yaml.load(open('budget-excel.yaml')) if os.path.exists('budget-excel.yaml') else {}

def walk(v, path):  # a path's parent and its last key: indices, or a list entry by name/item
    keys = path.split('.')
    for k in keys[:-1]: v = v[pick(v, k)]
    return v, pick(v, keys[-1])
def pick(v, k):
    if not isinstance(v, list): return k
    return int(k) if k.isdigit() else next(i for i, e in enumerate(v) if k in (e.get('name'), e.get('item')))
def get(v, path): p, k = walk(v, path); return p[k]
def cell(wb, ref): sheet, c = ref.split('!'); return wb[sheet][c]

if sys.argv[1:2] == ['import']:
    src = sys.argv[2]
    wb = load_workbook(src, data_only=True)  # the values Excel last computed
    b = yaml.load(open('budget.yaml'))
    for path, ref in (xl.get('inputs') or {}).items():
        p, k = walk(b, path); v = cell(wb, ref).value
        if v is not None: p[k] = v
    yaml.dump(b, open('budget.yaml', 'w'))
    c, off = computed(), []
    for path, ref in (xl.get('check') or {}).items():
        theirs = cell(wb, ref).value
        if isinstance(theirs, (int, float)) and abs(get(c, path) - theirs) > 1: off.append(f'{path}: workbook {theirs:,.2f}, tydig {get(c, path):,.2f}')
    print(f'budget.yaml updated from {src}' + (f"; totals disagree: {'; '.join(off)}" if off else '; totals agree' if xl.get('check') else ''))
    sys.exit(1 if off else 0)

os.makedirs('out', exist_ok=True)
if xl.get('template') and os.path.exists(xl['template']) and xl.get('inputs'):
    wb, b = load_workbook(xl['template']), yaml.load(open('budget.yaml'))
    for path, ref in xl['inputs'].items(): cell(wb, ref).value = get(b, path)
    wb.save('out/budget.xlsx'); print(f"out/budget.xlsx (budget.yaml's inputs in {xl['template']})"); sys.exit()
if xl: print(f"note: {xl.get('template')} or its inputs map is missing (budget-excel.yaml); a plain workbook instead")
b = computed()
wb = Workbook(); ws = wb.active; ws.title = 'Budget'
n = b['years']; col = lambda y: chr(ord('B') + y)
ws.append(['', *[f'Year {y + 1}' for y in range(n)], 'Total'])
def row(label, vals):
    ws.append([label, *[round(v, 2) for v in vals]]); r = ws.max_row
    ws.cell(r, n + 2, f'=SUM(B{r}:{col(n - 1)}{r})'); return r
first = ws.max_row + 1
for p in b['personnel']:
    row(f"{p['name']}: salary", p['salary']); row(f"{p['name']}: fringe", p['fringe'])
for o in b['other']: row(o['item'], o['cost'])
last = ws.max_row
ws.append(['Direct costs', *[f'=SUM({col(y)}{first}:{col(y)}{last})' for y in range(n + 1)]]); d = ws.max_row
row('F&A', b['fa'])
ws.append(['Total', *[f'={col(y)}{d}+{col(y)}{d + 1}' for y in range(n + 1)]])
row('NIH modules ($25,000 of direct costs)', b['modules'])
ws.column_dimensions['A'].width = 42
wb.save('out/budget.xlsx'); print('out/budget.xlsx')
